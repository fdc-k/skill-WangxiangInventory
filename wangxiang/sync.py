"""把差异写回 WooCommerce，并输出逐条结果日志。"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .console import arrow as _arrow
from .diff import S_CHANGE, STATUS_LABEL, _payload_for
from .woo import WooError, WooClient

FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FONT_HEAD = Font(color="FFFFFF", bold=True)
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_FAIL = PatternFill("solid", fgColor="FFC7CE")
FONT_OK = Font(color="006100")
FONT_FAIL = Font(color="9C0006")

LOG_COLUMNS = ["序号", "结果", "来源商品", "尺码", "本店 SKU", "本店商品", "变体 ID",
               "更新前", "更新请求", "更新后", "错误信息"]


def select_changes(diffs, *, only_sku=None, only_source=None):
    """按需缩小同步范围（例如先只同步 Sp5der 试水）。"""
    out = []
    for d in diffs:
        if d.status != S_CHANGE:
            continue
        if only_source:
            wanted = {x.strip().lower() for x in only_source if x.strip()}
            if (d.source or "").lower() not in wanted:
                continue
        if only_sku:
            needle = only_sku.strip().lower()
            hay = f"{d.parent_sku} {d.product_name} {d.source_product}".lower()
            if needle not in hay:
                continue
        out.append(d)
    return out


def write_snapshot(results, path, meta: dict | None = None):
    """把「改之前的原始值」存成快照，万一要回滚可以一键还原。"""
    import json
    entries = []
    for r in results:
        if r.get("结果") not in ("成功", "演练（未提交）"):
            continue
        entries.append({
            "product_id": r.get("_product_id"),
            "variation_id": r.get("变体 ID"),
            "sku": r.get("本店 SKU"),
            "before": r.get("_before"),
        })
    data = {"created_at": (meta or {}).get("generated_at", ""),
            "store": (meta or {}).get("store", ""),
            "dry_run": bool((meta or {}).get("dry_run")),
            "entries": entries}
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def rollback(snapshot_path, client: WooClient, *, dry_run=True, verbose=True):
    """按快照把库存还原回去。"""
    import json
    data = json.loads(Path(snapshot_path).read_text(encoding="utf-8"))
    entries = data.get("entries", [])
    results, ok, failed = [], 0, 0
    for i, e in enumerate(entries, 1):
        before = e.get("before") or {}
        payload = {}
        for k in ("manage_stock", "stock_quantity", "stock_status"):
            if k in before:
                payload[k] = before[k]
        if not payload or not e.get("variation_id"):
            continue
        entry = {"序号": i, "结果": "演练（未提交）" if dry_run else "",
                 "来源商品": "-", "尺码": "-", "本店 SKU": e.get("sku", ""),
                 "本店商品": "-", "变体 ID": e.get("variation_id"),
                 "更新前": "", "更新请求": json.dumps(payload, ensure_ascii=False),
                 "更新后": "", "错误信息": ""}
        if dry_run:
            results.append(entry)
            continue
        try:
            after = client.update_variation(e["product_id"], e["variation_id"], payload)
            entry["结果"] = "成功"
            entry["更新后"] = _fmt({k: (after or {}).get(k) for k in
                                    ("sku", "manage_stock", "stock_quantity", "stock_status")})
            ok += 1
        except Exception as exc:  # noqa: BLE001
            entry["结果"] = "失败"
            entry["错误信息"] = str(exc)[:400]
            failed += 1
        results.append(entry)
        if verbose:
            print(f"  [{i}/{len(entries)}] {entry['结果']} {e.get('sku')} -> {entry['更新请求']}")
    return results, {"total": len(entries), "success": ok, "failed": failed,
                     "skipped": len(entries) - ok - failed, "dry_run": dry_run}


def apply_updates(diffs, catalog, cfg, client: WooClient, *, dry_run=True, verbose=True,
                  subset=None):
    """执行更新。``subset`` 为空时对所有「有差异」条目动手。返回 (结果列表, 统计)。"""
    sync_cfg = cfg.get("sync", {})
    changes = subset if subset is not None else [d for d in diffs if d.status == S_CHANGE]
    results = []
    ok = failed = skipped = 0

    for i, d in enumerate(changes, 1):
        variation = None
        for v in catalog.get("variations", {}).get(str(d.product_id), []):
            if v.get("id") == d.variation_id:
                variation = v
                break
        if variation is None:
            results.append({"序号": i, "结果": "跳过", "来源商品": d.source_product,
                            "尺码": d.size, "本店 SKU": d.parent_sku, "本店商品": d.product_name,
                            "变体 ID": d.variation_id, "更新前": "", "更新请求": "",
                            "更新后": "", "错误信息": "找不到变体"})
            skipped += 1
            continue

        payload = _payload_for(sync_cfg, d, variation)
        before = {k: variation.get(k) for k in ("sku", "manage_stock", "stock_quantity", "stock_status")}
        entry = {
            "序号": i, "结果": "演练（未提交）" if dry_run else "",
            "来源商品": d.source_product, "尺码": d.size,
            "本店 SKU": d.parent_sku, "本店商品": d.product_name,
            "变体 ID": d.variation_id,
            "更新前": _fmt(before), "更新请求": json.dumps(payload, ensure_ascii=False),
            "更新后": "", "错误信息": "",
            "_product_id": d.product_id, "_before": before,
        }
        if dry_run:
            results.append(entry)
            skipped += 1
            continue

        try:
            after_raw = client.update_variation(d.product_id, d.variation_id, payload)
            after = {k: (after_raw or {}).get(k) for k in
                     ("sku", "manage_stock", "stock_quantity", "stock_status")}
            entry["结果"] = "成功"
            entry["更新后"] = _fmt(after)
            ok += 1
        except (WooError, Exception) as exc:      # noqa: BLE001 - 逐条记录，不中断整批
            entry["结果"] = "失败"
            entry["错误信息"] = str(exc)[:400]
            failed += 1
        results.append(entry)
        if verbose:
            print(f"  [{i}/{len(changes)}] {entry['结果']} {d.source_product} / {d.size} "
                  f"-> {entry['更新请求']}")

    stats = {"total": len(changes), "success": ok, "failed": failed, "skipped": skipped,
             "dry_run": dry_run}
    return results, stats


def _fmt(d: dict) -> str:
    if not d:
        return ""
    return (f"数量={d.get('stock_quantity')} 状态="
            f"{STATUS_LABEL.get(d.get('stock_status'), d.get('stock_status'))} "
            f"管理库存={'是' if d.get('manage_stock') else '否'}")


def write_sync_log(results, stats, path, meta: dict | None = None):
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "同步结果"
    ws.append(LOG_COLUMNS)
    for c in range(1, len(LOG_COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = FILL_HEAD
        cell.font = FONT_HEAD
        cell.alignment = Alignment(horizontal="center")
    widths = [6, 14, 32, 8, 20, 32, 10, 32, 34, 32, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in results:
        ws.append([r.get(c, "") for c in LOG_COLUMNS])
        row = ws.max_row
        if r.get("结果") == "成功":
            for c in range(1, len(LOG_COLUMNS) + 1):
                ws.cell(row=row, column=c).fill = FILL_OK
                ws.cell(row=row, column=c).font = FONT_OK
        elif r.get("结果") == "失败":
            for c in range(1, len(LOG_COLUMNS) + 1):
                ws.cell(row=row, column=c).fill = FILL_FAIL
                ws.cell(row=row, column=c).font = FONT_FAIL
    ws.freeze_panes = "A2"

    # 按商品汇总：一个商品一行，一眼看出「哪些商品成功了、哪些没有」
    ws_p = wb.create_sheet("按商品汇总")
    head_p = ["本店 SKU", "本店商品", "更新条数", "成功", "失败", "改变内容"]
    ws_p.append(head_p)
    for c in range(1, len(head_p) + 1):
        cell = ws_p.cell(row=1, column=c)
        cell.fill = FILL_HEAD
        cell.font = FONT_HEAD
        cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate([22, 38, 10, 8, 8, 80], 1):
        ws_p.column_dimensions[get_column_letter(i)].width = w
    grouped = {}
    for r in results:
        key = (r.get("本店 SKU", ""), r.get("本店商品", ""))
        g = grouped.setdefault(key, {"n": 0, "ok": 0, "fail": 0, "bits": []})
        g["n"] += 1
        if r.get("结果") == "成功":
            g["ok"] += 1
        elif r.get("结果") == "失败":
            g["fail"] += 1
        g["bits"].append(f'{r.get("尺码")}: {r.get("更新前")} {_arrow()} {r.get("更新请求")}')
    for (sku, name), g in grouped.items():
        ws_p.append([sku, name, g["n"], g["ok"], g["fail"], "；".join(g["bits"])])
        row = ws_p.max_row
        fill = FILL_OK if g["fail"] == 0 and g["ok"] else (FILL_FAIL if g["fail"] else None)
        if fill:
            for c in range(1, len(head_p) + 1):
                ws_p.cell(row=row, column=c).fill = fill
    ws_p.freeze_panes = "A2"

    ws2 = wb.create_sheet("汇总")
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 60
    info = [
        ("执行时间", (meta or {}).get("generated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))),
        ("店铺", (meta or {}).get("store", "")),
        ("模式", "演练（未提交任何改动）" if stats.get("dry_run") else "真实更新"),
        ("计划更新条目", stats.get("total")),
        ("成功", stats.get("success")),
        ("失败", stats.get("failed")),
        ("跳过/演练", stats.get("skipped")),
    ]
    for i, (k, v) in enumerate(info, start=1):
        ws2.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws2.cell(row=i, column=2, value=v)
    wb.save(path)
    return path
