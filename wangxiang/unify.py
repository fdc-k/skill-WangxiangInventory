"""库存文件 -> 统一格式 Excel。"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as cfgmod
from .console import warn as _warn
from .extract_excel import extract_excel
from .extract_pdf import extract_pdf_pallet
from .models import UNIFIED_COLUMNS, InventoryRow

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=12)


def extract_one_source(source: dict, cfg: dict):
    """按配置里的 adapter 调用对应解析器。"""
    adapter = source.get("adapter", "excel_table")
    path = cfgmod.resolve_path(cfg, source["file"])
    if not path.exists():
        raise FileNotFoundError(f"来源文件不存在：{path}")
    options = source.get("options", {}) or {}
    block_detail = None
    if adapter == "excel_table":
        rows, stats = extract_excel(path, source["id"], options)
    elif adapter == "pdf_pallet":
        rows, stats, block_detail = extract_pdf_pallet(path, source["id"], options)
    else:
        raise ValueError(f"未知的 adapter：{adapter}")
    stats["label"] = source.get("label", source["id"])
    return rows, stats, block_detail


def write_unified(rows: list[InventoryRow], path: Path, title: str = ""):
    wb = Workbook()
    ws = wb.active
    ws.title = "统一库存"
    headers = list(UNIFIED_COLUMNS)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        d = row.as_dict()
        ws.append([d.get(h) if d.get(h) is not None else "" for h in headers])
    widths = {"source": 12, "category": 18, "source_product": 40, "size": 8,
              "quantity": 10, "availability": 13, "source_ref": 46}
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def run_extract(cfg: dict, only: list[str] | None = None, verbose: bool = True) -> dict:
    """抽取所有来源，写出统一格式文件。返回统计信息。"""
    out_dir = cfgmod.output_dir(cfg) / "unified"
    out_dir.mkdir(parents=True, exist_ok=True)

    results, all_rows, report = {}, [], []
    for source in cfg.get("sources", []):
        sid = source["id"]
        if only and sid not in only:
            continue
        rows, stats, blocks = extract_one_source(source, cfg)
        all_rows.extend(rows)
        fname = f"{sid}_统一库存.xlsx"
        write_unified(rows, out_dir / fname, stats.get("label", sid))
        results[sid] = {"rows": rows, "stats": stats, "blocks": blocks,
                        "file": str(out_dir / fname)}
        report.append(stats)
        if verbose:
            print(f"  [{sid}] {stats.get('label','')}：{stats['rows']} 行 / "
                  f"{stats['products']} 个商品 -> {fname}")
            if stats.get("orphan_skus"):
                print(f"      {_warn()}未匹配到商品块的 SKU：{stats['orphan_skus']}")
    # 只有「处理全部来源」时才写合并文件，避免 --sources 局部跑覆盖掉完整的合并表
    if all_rows and not only:
        write_unified(all_rows, out_dir / "全部库存_合并.xlsx", "合并")
    (out_dir / "extract_stats.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"sources": results, "rows": all_rows, "dir": str(out_dir)}
